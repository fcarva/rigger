"""Leitura por visão: transcreve slides em imagem (PPTX) para texto técnico.

Útil para apostilas em que cada slide é apenas uma imagem (sem texto extraível).
O resultado é salvo em um arquivo `.transcricao.md` reutilizável, para que a
transcrição (que consome tokens de visão) seja paga apenas uma vez.
"""

from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path

from .cliente_claude import ClienteClaude
from .config import Config
from .prompts import INSTRUCOES_TRANSCRICAO

# shape_type 13 = PICTURE (python-pptx MSO_SHAPE_TYPE.PICTURE)
_TIPO_IMAGEM = 13


def _redimensionar(dados: bytes, media_type: str, max_lado: int) -> tuple[bytes, str]:
    """Reduz a imagem para no máximo `max_lado` px (lado maior), como JPEG.

    Controla o custo de visão. Se o Pillow não estiver disponível, devolve o original.
    """
    try:
        from PIL import Image
    except ImportError:
        return dados, media_type

    try:
        imagem = Image.open(BytesIO(dados))
        if max(imagem.size) > max_lado:
            imagem.thumbnail((max_lado, max_lado))
        if imagem.mode not in ("RGB", "L"):
            imagem = imagem.convert("RGB")
        buffer = BytesIO()
        imagem.save(buffer, format="JPEG", quality=82)
        return buffer.getvalue(), "image/jpeg"
    except Exception:  # noqa: BLE001 - em caso de falha, usa o original
        return dados, media_type


def extrair_imagens_pptx(
    caminho: Path, max_lado: int = 1500
) -> list[tuple[int, str, str]]:
    """Extrai a(s) imagem(ns) de cada slide do PPTX.

    Devolve lista de tuplas (numero_do_slide, dados_base64, media_type).
    """
    from pptx import Presentation

    apresentacao = Presentation(str(caminho))
    imagens: list[tuple[int, str, str]] = []
    for numero, slide in enumerate(apresentacao.slides, start=1):
        for forma in slide.shapes:
            if forma.shape_type == _TIPO_IMAGEM:
                imagem = forma.image
                dados, media = _redimensionar(imagem.blob, imagem.content_type, max_lado)
                imagens.append((numero, base64.b64encode(dados).decode(), media))
    return imagens


def transcrever_pptx(
    caminho: Path,
    *,
    lote: int = 6,
    destino: Path | None = None,
    config: Config | None = None,
    cliente: ClienteClaude | None = None,
    progresso=None,
) -> tuple[str, Path]:
    """Transcreve um PPTX de slides-imagem para Markdown e salva em arquivo.

    `lote` = quantos slides por chamada à API. `progresso` (opcional) é um callable
    chamado a cada lote com (slides_feitos, total). Devolve (texto, caminho_salvo).
    """
    config = config or Config()
    cliente = cliente or ClienteClaude(config)

    imagens = extrair_imagens_pptx(caminho)
    if not imagens:
        raise ValueError(f"Nenhuma imagem encontrada em {caminho.name}.")

    partes: list[str] = []
    total = len(imagens)
    for inicio in range(0, total, lote):
        bloco = imagens[inicio : inicio + lote]
        partes.append(cliente.transcrever_imagens(bloco, INSTRUCOES_TRANSCRICAO))
        if progresso is not None:
            progresso(min(inicio + lote, total), total)

    texto = "\n\n".join(partes)

    if destino is None:
        destino = caminho.with_suffix(".transcricao.md")
    destino.write_text(texto, encoding="utf-8")
    return texto, destino
