"""Monta um documento Word (.docx) a partir de uma `ApreciacaoRisco`."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

from ..modelos import ApreciacaoRisco, ChecklistInspecao, ItemRisco, PlanoIcamento

# Cores de fundo (hex) por classificação de risco.
_CORES_RISCO = {
    "Baixo": "C6EFCE",
    "Médio": "FFEB9C",
    "Alto": "FFC7A0",
    "Crítico": "FF7C80",
}


def _sombrear_celula(celula, cor_hex: str) -> None:
    """Define a cor de fundo de uma célula da tabela."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    celula._tc.get_or_add_tcPr().append(shd)


def _linha_info(documento, rotulo: str, valor: str) -> None:
    paragrafo = documento.add_paragraph()
    run = paragrafo.add_run(f"{rotulo}: ")
    run.bold = True
    paragrafo.add_run(valor)


def gerar_docx(apr: ApreciacaoRisco, caminho_saida: Path) -> Path:
    """Gera o .docx da apreciação de risco e devolve o caminho do arquivo."""
    documento = Document()

    titulo = documento.add_heading("Apreciação de Risco — Içamento de Carga", level=0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Cabeçalho de identificação
    _linha_info(documento, "Atividade", apr.atividade)
    if apr.local:
        _linha_info(documento, "Local", apr.local)
    if apr.equipe:
        _linha_info(documento, "Equipe", ", ".join(apr.equipe))
    if apr.normas_referencia:
        _linha_info(documento, "Normas de referência", ", ".join(apr.normas_referencia))

    documento.add_paragraph()

    # Tabela de riscos
    colunas = [
        "Etapa",
        "Perigo",
        "Danos possíveis",
        "P",
        "S",
        "Risco",
        "Classificação",
        "Medidas de controle",
        "Responsável",
    ]
    tabela = documento.add_table(rows=1, cols=len(colunas))
    tabela.style = "Light Grid Accent 1"

    for celula, texto in zip(tabela.rows[0].cells, colunas):
        celula.paragraphs[0].add_run(texto).bold = True

    for item in apr.itens:
        celulas = tabela.add_row().cells
        celulas[0].text = item.etapa
        celulas[1].text = item.perigo
        celulas[2].text = "\n".join(f"• {d}" for d in item.danos_possiveis)
        celulas[3].text = str(item.probabilidade)
        celulas[4].text = str(item.severidade)
        celulas[5].text = str(item.risco)
        celulas[6].text = item.classificacao
        celulas[7].text = "\n".join(f"• {m}" for m in item.medidas_controle)
        celulas[8].text = item.responsavel
        _sombrear_celula(celulas[6], _CORES_RISCO.get(item.classificacao, "FFFFFF"))

    # Recomendações gerais
    if apr.recomendacoes_gerais:
        documento.add_heading("Recomendações gerais", level=1)
        for rec in apr.recomendacoes_gerais:
            documento.add_paragraph(rec, style="List Bullet")

    # Legenda da matriz de risco
    documento.add_heading("Legenda — Matriz de risco (P × S, escala 1–5)", level=2)
    legenda = (
        "Baixo (1–4) • Médio (5–9) • Alto (10–14) • Crítico (15–25). "
        "P = probabilidade, S = severidade."
    )
    documento.add_paragraph(legenda).runs[0].font.size = Pt(9)

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    documento.save(str(caminho_saida))
    return caminho_saida


# ---------------------------------------------------------------------------
# Plano de içamento (rigging plan)
# ---------------------------------------------------------------------------


def _bloco_assinaturas(documento, papeis: list[str]) -> None:
    """Adiciona linhas de assinatura (ex.: Elaborado / Verificado / Aprovado)."""
    documento.add_paragraph()
    for papel in papeis:
        documento.add_paragraph("_" * 40)
        documento.add_paragraph(f"{papel} — nome, assinatura e data")


def gerar_docx_plano(plano: PlanoIcamento, caminho_saida: Path) -> Path:
    """Gera o .docx do plano de içamento e devolve o caminho do arquivo."""
    documento = Document()
    titulo = documento.add_heading("Plano de Içamento (Rigging Plan)", level=0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if plano.obra:
        _linha_info(documento, "Obra/local", plano.obra)

    documento.add_heading("Carga", level=1)
    _linha_info(documento, "Descrição", plano.descricao_carga)
    _linha_info(documento, "Peso", plano.peso_carga)
    if plano.dimensoes_carga:
        _linha_info(documento, "Dimensões", plano.dimensoes_carga)
    if plano.centro_gravidade:
        _linha_info(documento, "Centro de gravidade", plano.centro_gravidade)

    documento.add_heading("Equipamento", level=1)
    _linha_info(documento, "Equipamento", plano.equipamento)
    for rotulo, valor in [
        ("Capacidade", plano.capacidade_equipamento),
        ("Raio de operação", plano.raio_operacao),
        ("Comprimento de lança", plano.comprimento_lanca),
        ("Utilização da tabela de carga", plano.percentual_utilizacao),
    ]:
        if valor:
            _linha_info(documento, rotulo, valor)

    if plano.acessorios:
        documento.add_heading("Acessórios de içamento", level=1)
        tabela = documento.add_table(rows=1, cols=4)
        tabela.style = "Light Grid Accent 1"
        for celula, texto in zip(tabela.rows[0].cells, ["Tipo", "Capacidade (WLL)", "Qtde", "Observação"]):
            celula.paragraphs[0].add_run(texto).bold = True
        for acessorio in plano.acessorios:
            celulas = tabela.add_row().cells
            celulas[0].text = acessorio.tipo
            celulas[1].text = acessorio.capacidade
            celulas[2].text = acessorio.quantidade
            celulas[3].text = acessorio.observacao

    if plano.sequencia_operacao:
        documento.add_heading("Sequência da operação", level=1)
        for passo in plano.sequencia_operacao:
            documento.add_paragraph(passo, style="List Number")

    if plano.riscos_criticos:
        documento.add_heading("Riscos críticos", level=1)
        for risco in plano.riscos_criticos:
            documento.add_paragraph(risco, style="List Bullet")

    if plano.criterios_seguranca:
        documento.add_heading("Critérios e limites de segurança", level=1)
        for criterio in plano.criterios_seguranca:
            documento.add_paragraph(criterio, style="List Bullet")

    if plano.responsaveis:
        documento.add_heading("Responsáveis", level=1)
        tabela = documento.add_table(rows=1, cols=2)
        tabela.style = "Light Grid Accent 1"
        for celula, texto in zip(tabela.rows[0].cells, ["Função", "Atribuição"]):
            celula.paragraphs[0].add_run(texto).bold = True
        for resp in plano.responsaveis:
            celulas = tabela.add_row().cells
            celulas[0].text = resp.funcao
            celulas[1].text = resp.atribuicao

    if plano.normas_referencia:
        documento.add_heading("Normas de referência", level=1)
        documento.add_paragraph(", ".join(plano.normas_referencia))

    _bloco_assinaturas(documento, ["Elaborado por", "Verificado por", "Aprovado por"])

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    documento.save(str(caminho_saida))
    return caminho_saida


# ---------------------------------------------------------------------------
# Checklist de inspeção de acessórios
# ---------------------------------------------------------------------------


def gerar_docx_checklist(checklist: ChecklistInspecao, caminho_saida: Path) -> Path:
    """Gera o .docx do checklist de inspeção e devolve o caminho do arquivo."""
    documento = Document()
    titulo = documento.add_heading(f"Checklist de Inspeção — {checklist.acessorio}", level=0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if checklist.periodicidade:
        _linha_info(documento, "Periodicidade", checklist.periodicidade)
    if checklist.normas_referencia:
        _linha_info(documento, "Normas de referência", ", ".join(checklist.normas_referencia))

    documento.add_paragraph()

    colunas = ["#", "Item de inspeção", "Critério de aprovação", "Conforme? (C/NC)", "Referência"]
    tabela = documento.add_table(rows=1, cols=len(colunas))
    tabela.style = "Light Grid Accent 1"
    for celula, texto in zip(tabela.rows[0].cells, colunas):
        celula.paragraphs[0].add_run(texto).bold = True

    for indice, item in enumerate(checklist.itens, start=1):
        celulas = tabela.add_row().cells
        celulas[0].text = str(indice)
        celulas[1].text = item.item
        celulas[2].text = item.criterio
        celulas[3].text = ""  # preenchido na inspeção
        celulas[4].text = item.referencia

    if checklist.criterios_descarte:
        documento.add_heading("Critérios de descarte", level=1)
        for criterio in checklist.criterios_descarte:
            documento.add_paragraph(criterio, style="List Bullet")

    _bloco_assinaturas(documento, ["Inspecionado por"])

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    documento.save(str(caminho_saida))
    return caminho_saida


__all__ = ["gerar_docx", "gerar_docx_plano", "gerar_docx_checklist", "ItemRisco"]
