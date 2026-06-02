"""Monta um arquivo PowerPoint (.pptx) a partir de um `Treinamento`."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

from ..modelos import Treinamento

# Cor de destaque (azul corporativo sóbrio).
from pptx.dml.color import RGBColor

_AZUL = RGBColor(0x1F, 0x4E, 0x79)


def _slide_titulo(prs: Presentation, titulo: str, subtitulo: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = titulo
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitulo


def _slide_secao(prs: Presentation, titulo: str) -> None:
    layout = prs.slide_layouts[5]  # somente título
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = titulo
    try:
        slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = _AZUL
    except (AttributeError, IndexError):
        pass


def _slide_bullets(
    prs: Presentation, titulo: str, bullets: list[str], notas: str = ""
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # título + conteúdo
    slide.shapes.title.text = titulo

    corpo = slide.placeholders[1].text_frame
    corpo.clear()
    for i, item in enumerate(bullets):
        paragrafo = corpo.paragraphs[0] if i == 0 else corpo.add_paragraph()
        paragrafo.text = item
        paragrafo.font.size = Pt(20)

    if notas:
        slide.notes_slide.notes_text_frame.text = notas


def gerar_pptx(treinamento: Treinamento, caminho_saida: Path) -> Path:
    """Gera o .pptx do treinamento e devolve o caminho do arquivo."""
    prs = Presentation()

    # Capa
    _slide_titulo(
        prs,
        treinamento.titulo,
        f"{treinamento.publico_alvo}  •  Carga horária: {treinamento.carga_horaria}",
    )

    # Objetivos
    if treinamento.objetivos:
        _slide_bullets(prs, "Objetivos do treinamento", treinamento.objetivos)

    # Pré-requisitos
    if treinamento.pre_requisitos:
        _slide_bullets(prs, "Pré-requisitos", treinamento.pre_requisitos)

    # Módulos
    for indice, modulo in enumerate(treinamento.modulos, start=1):
        _slide_secao(prs, f"Módulo {indice} — {modulo.titulo}")
        if modulo.objetivo:
            _slide_bullets(prs, "Objetivo do módulo", [modulo.objetivo])
        for slide in modulo.slides:
            _slide_bullets(prs, slide.titulo, slide.topicos, slide.notas_apresentador)

    # Normas de referência
    if treinamento.normas_referencia:
        _slide_bullets(prs, "Normas de referência", treinamento.normas_referencia)

    # Avaliação
    if treinamento.avaliacao:
        _slide_bullets(prs, "Avaliação de aprendizagem", treinamento.avaliacao)

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(caminho_saida))
    return caminho_saida
