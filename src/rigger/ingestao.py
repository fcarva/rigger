"""Leitura de materiais de curso e normas da base de conhecimento.

Formatos suportados: .md, .txt, .pdf, .docx, .pptx
"""

from __future__ import annotations

from pathlib import Path

EXTENSOES_SUPORTADAS = {".md", ".txt", ".pdf", ".docx", ".pptx"}


def _ler_pdf(caminho: Path) -> str:
    from pypdf import PdfReader

    leitor = PdfReader(str(caminho))
    paginas = [(pagina.extract_text() or "") for pagina in leitor.pages]
    return "\n".join(paginas)


def _ler_docx(caminho: Path) -> str:
    from docx import Document

    documento = Document(str(caminho))
    return "\n".join(paragrafo.text for paragrafo in documento.paragraphs)


def _texto_de_forma(forma) -> list[str]:
    """Extrai texto de uma forma do slide, recursivamente.

    Cobre caixas de texto, placeholders, tabelas e grupos (que contêm outras formas).
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    textos: list[str] = []

    # Grupos: percorre as formas-filhas.
    if forma.shape_type == MSO_SHAPE_TYPE.GROUP:
        for filha in forma.shapes:
            textos.extend(_texto_de_forma(filha))
        return textos

    # Tabelas: lê célula a célula, linha por linha.
    if forma.has_table:
        for linha in forma.table.rows:
            celulas = [celula.text.strip() for celula in linha.cells]
            if any(celulas):
                textos.append(" | ".join(celulas))
        return textos

    # Caixas de texto / placeholders.
    if forma.has_text_frame:
        conteudo = forma.text_frame.text.strip()
        if conteudo:
            textos.append(conteudo)

    return textos


def _ler_pptx(caminho: Path) -> str:
    from pptx import Presentation

    apresentacao = Presentation(str(caminho))
    partes: list[str] = []
    for numero, slide in enumerate(apresentacao.slides, start=1):
        textos: list[str] = []
        for forma in slide.shapes:
            textos.extend(_texto_de_forma(forma))
        # Notas do apresentador, quando houver.
        if slide.has_notes_slide:
            nota = slide.notes_slide.notes_text_frame.text.strip()
            if nota:
                textos.append(f"[Notas do slide] {nota}")
        if textos:
            partes.append(f"[Slide {numero}]\n" + "\n".join(textos))
    return "\n\n".join(partes)


def ler_arquivo(caminho: Path) -> str:
    """Lê um único arquivo e devolve seu texto."""
    sufixo = caminho.suffix.lower()
    if sufixo in {".md", ".txt"}:
        return caminho.read_text(encoding="utf-8", errors="ignore")
    if sufixo == ".pdf":
        return _ler_pdf(caminho)
    if sufixo == ".docx":
        return _ler_docx(caminho)
    if sufixo == ".pptx":
        return _ler_pptx(caminho)
    raise ValueError(f"Formato não suportado: {caminho.name}")


def ler_caminho(caminho: Path) -> str:
    """Lê um arquivo ou todos os arquivos suportados de um diretório (recursivo).

    Cada arquivo é precedido por um cabeçalho com seu nome, para o modelo saber a origem.
    """
    if caminho.is_file():
        return f"### Arquivo: {caminho.name}\n\n{ler_arquivo(caminho)}"

    if not caminho.exists():
        raise FileNotFoundError(f"Caminho não encontrado: {caminho}")

    partes: list[str] = []
    for arquivo in sorted(caminho.rglob("*")):
        if arquivo.is_file() and arquivo.suffix.lower() in EXTENSOES_SUPORTADAS:
            try:
                texto = ler_arquivo(arquivo)
            except Exception as erro:  # noqa: BLE001 - registra e segue
                texto = f"[Falha ao ler {arquivo.name}: {erro}]"
            partes.append(f"### Arquivo: {arquivo.relative_to(caminho)}\n\n{texto}")

    if not partes:
        raise FileNotFoundError(
            f"Nenhum arquivo suportado ({', '.join(sorted(EXTENSOES_SUPORTADAS))}) "
            f"encontrado em: {caminho}"
        )
    return "\n\n---\n\n".join(partes)


def ler_normas(dir_normas: Path) -> str:
    """Lê toda a pasta de normas. Devolve string vazia se a pasta não existir/estiver vazia."""
    try:
        return ler_caminho(dir_normas)
    except FileNotFoundError:
        return ""
