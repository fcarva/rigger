"""Testes da ingestão de PPTX (extração de texto de tabelas), sem API."""

from __future__ import annotations

from pathlib import Path


def test_ler_pptx_tabela(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    from rigger.ingestao import ler_arquivo

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    tabela = slide.shapes.add_table(
        2, 2, Inches(1), Inches(1), Inches(4), Inches(1)
    ).table
    tabela.cell(0, 0).text = "Acessório"
    tabela.cell(0, 1).text = "Capacidade"
    tabela.cell(1, 0).text = "Cinta"
    tabela.cell(1, 1).text = "2 t"

    destino = tmp_path / "tabela.pptx"
    prs.save(str(destino))

    texto = ler_arquivo(destino)
    assert "Acessório | Capacidade" in texto
    assert "Cinta | 2 t" in texto
